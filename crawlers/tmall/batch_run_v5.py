#!/usr/bin/env python3
"""
帝泊洱评论爬虫 v5 — 适配天猫SSR新版Drawer评论区
关键：评论区在左侧弹出的Drawer面板中，需滚动Drawer内的scroll容器
"""
import sys, os, time, re

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "output")
PROFILE_DIR = os.path.join(SCRIPT_DIR, "browser_profile")
os.makedirs(OUTPUT_DIR, exist_ok=True)

from playwright.sync_api import sync_playwright

PRODUCTS = [
    ("🎁 试饮装(顺手买)", "https://e.tb.cn/h.izPFxG9RD4SnZ89"),
    ("🍵 甘醇陈皮菊花糯香40支", "https://e.tb.cn/h.iBhRLAptqLL5o9j"),
    ("🌾 糯香10支装", "https://e.tb.cn/h.iyn1AmS5xWvanJz"),
    ("📦 100支量贩装", "https://e.tb.cn/h.iynXx7oUIH6zOd3"),
    ("🌹 重瓣玫瑰普洱茶珍", "https://e.tb.cn/h.iCDjYRN1AV1ZfjL"),
    ("🌸 玫瑰菊花60支", "https://e.tb.cn/h.iynXCURiTiWuer7"),
    ("🍊 陈皮小青柑味10支", "https://e.tb.cn/h.iCDj8XmFbL3BhiD"),
    ("☕ 经典甘醇盒装10支", "https://e.tb.cn/h.iCD9vanubigOOmd"),
    ("🌼 茉莉花茶10支装", "https://e.tb.cn/h.izPyyBTOi0CArOE"),
    ("🆕 U先5支装", "https://e.tb.cn/h.iyn3HzoqyteneC5"),
]
MAX_PER_PRODUCT = 150
LOGIN_TIMEOUT = 180

def log(msg): print(f"  {msg}")

def create_browser(p):
    return p.chromium.launch_persistent_context(
        PROFILE_DIR, headless=False,
        args=["--disable-blink-features=AutomationControlled", "--no-sandbox"],
        viewport={"width": 1440, "height": 1000},
        locale="zh-CN",
    )

def wait_login(page, url):
    page.goto(url, timeout=60000, wait_until="domcontentloaded")
    page.wait_for_timeout(5000)
    login_kw = ["login.taobao.com", "login.tmall.com", "havanaone/login"]
    if any(k in page.url.lower() for k in login_kw):
        log(f"🔑 请在浏览器中登录淘宝（等待 {LOGIN_TIMEOUT} 秒）...")
        for _ in range(LOGIN_TIMEOUT // 2):
            page.wait_for_timeout(2000)
            if not any(k in page.url.lower() for k in login_kw):
                log("✅ 已登录")
                page.wait_for_timeout(3000)
                return True
        log("⏰ 登录超时")
        return False
    return True

def open_review_drawer(page):
    """打开评论Drawer面板"""
    # 方法1: 点"查看全部评价"按钮
    try:
        btn = page.locator('[class*="ShowButton"]').first
        if btn.is_visible(timeout=3000):
            btn.click()
            page.wait_for_timeout(2000)
            log("  ✅ 点击了「查看全部评价」")
            return True
    except Exception as e:
        log(f"  ⚠️ 点击ShowButton失败: {e}")
    
    # 方法2: 点评价Tab
    try:
        tab = page.locator('span:has-text("用户评价"), span:has-text("评价")').first
        if tab.is_visible(timeout=2000):
            tab.click()
            page.wait_for_timeout(2000)
            log("  ✅ 点击了评价Tab")
            return True
    except Exception as e:
        log(f"  ⚠️ 点击评价Tab失败: {e}")
    
    log("  ⚠️ 未打开评论Drawer")
    return False

def scroll_drawer_comments(page):
    """在Drawer面板内滚动评论列表"""
    log("  📜 在Drawer内滚动加载评论...")
    
    last_count = 0
    no_change = 0
    max_rounds = 60
    
    for rnd in range(max_rounds):
        # 在Drawer面板内滚动：找到可滚动的comments容器
        page.evaluate("""
        () => {
            // Drawer里的评论容器有 overflow-y: scroll
            const containers = document.querySelectorAll('[class*="comments--"]');
            for (const c of containers) {
                const style = window.getComputedStyle(c);
                if (style.overflowY === 'scroll' || style.overflowY === 'auto') {
                    c.scrollTop = c.scrollHeight;
                    break;
                }
            }
        }
        """)
        page.wait_for_timeout(1000)
        
        # 统计Drawer内的评论数
        count = page.evaluate("""
        () => {
            // 只数Drawer内的Comment
            const drawer = document.querySelector('[class*="Drawer--"]');
            if (drawer) {
                return drawer.querySelectorAll('[class*="Comment--"]').length;
            }
            // 备用：数overflow-y:scroll容器内的
            const containers = document.querySelectorAll('[class*="comments--"]');
            for (const c of containers) {
                const style = window.getComputedStyle(c);
                if (style.overflowY === 'scroll' || style.overflowY === 'auto') {
                    return c.querySelectorAll('[class*="Comment--"]').length;
                }
            }
            return document.querySelectorAll('[class*="Comment--"]').length;
        }
        """)
        
        if count == last_count:
            no_change += 1
            if no_change >= 8:
                break
        else:
            no_change = 0
            if rnd % 3 == 0:
                log(f"   ... {count} 条")
        last_count = count
        
        if count >= MAX_PER_PRODUCT:
            log(f"  ✓ 已达上限 {MAX_PER_PRODUCT} 条")
            break
    
    log(f"  ✓ 共 {count} 条评论")
    return count

def extract_reviews(page):
    """提取Drawer内评论"""
    js = """
    () => {
        // 优先从Drawer里取评论
        let comments = [];
        const drawer = document.querySelector('[class*="Drawer--"]');
        if (drawer) {
            comments = drawer.querySelectorAll('[class*="Comment--"]');
        }
        if (comments.length === 0) {
            // 备用：取所有Comment
            comments = document.querySelectorAll('[class*="Comment--"]');
        }
        
        const results = [];
        const seen = new Set();
        
        comments.forEach(el => {
            try {
                const nickEl = el.querySelector('[class*="userName--"]');
                const nick = nickEl ? nickEl.textContent.trim().slice(0, 20) : '';
                
                const metaEl = el.querySelector('[class*="meta--"]');
                let date = '', sku = '';
                if (metaEl) {
                    const t = metaEl.textContent.trim();
                    const dm = t.match(/(\\d{4}[-.\\/]\\d{2}[-.\\/]\\d{2})/);
                    if (dm) date = dm[1];
                    const si = t.indexOf('已购');
                    if (si >= 0) sku = t.substring(si + 3).trim();
                }
                
                const contentEl = el.querySelector('[class*="content--"]');
                let content = contentEl ? contentEl.textContent.trim() : '';
                if (!content) {
                    const allText = el.textContent || '';
                    const headerEl = el.querySelector('[class*="header--"]');
                    const hText = headerEl ? headerEl.textContent.trim() : '';
                    content = allText.replace(hText, '').replace(/商家回复[\\s\\S]*$/, '').trim();
                }
                
                const imgCount = el.querySelectorAll('img').length;
                const key = (nick + content.slice(0, 30)).replace(/\\s/g, '');
                if (content && content.length > 2 && !seen.has(key)) {
                    seen.add(key);
                    results.push({
                        nick, content: content.slice(0, 500), sku, date,
                        pics: String(imgCount),
                    });
                }
            } catch(e) {}
        });
        return results;
    }
    """
    return page.evaluate(js)

def scrape_product(page, name, url):
    log(f"\n{'─'*60}")
    log(f"  [{name}]  {url}")
    log(f"{'─'*60}")
    
    page.goto(url, timeout=60000, wait_until="domcontentloaded")
    page.wait_for_timeout(6000)
    
    # 检查登录，如果跳转到登录页就等待用户登录
    login_kw = ["login.taobao.com", "login.tmall.com", "havanaone/login", "login_jump"]
    if any(k in page.url.lower() for k in login_kw):
        log(f"🔑 需要登录淘宝（等待 {LOGIN_TIMEOUT} 秒）...")
        for _ in range(LOGIN_TIMEOUT // 2):
            page.wait_for_timeout(2000)
            if not any(k in page.url.lower() for k in login_kw):
                log("✅ 已登录")
                page.wait_for_timeout(3000)
                break
        else:
            log("⏰ 登录超时")
    
    log(f"  → {page.url[:100]}...")
    
    # 打开评价Drawer
    open_review_drawer(page)
    page.wait_for_timeout(2000)
    
    # 在Drawer里滚动
    scroll_drawer_comments(page)
    
    # 提取
    reviews = extract_reviews(page)
    
    # 去重
    seen, unique = set(), []
    for r in reviews:
        k = (r.get("nick",""), r.get("content","")[:40])
        if k not in seen:
            seen.add(k)
            unique.append(r)
    
    log(f"  ✅ {len(unique)} 条有效评论")
    return unique

def save_excel(reviews, name, pid, is_merged=False, all_reviews=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    
    ts = time.strftime("%Y%m%d_%H%M%S")
    
    if is_merged:
        filepath = os.path.join(OUTPUT_DIR, f"reviews_帝泊洱合集_{ts}.xlsx")
        ws_title = "全部评论"
        headers = ["序号","来源商品","商品名称","商品ID","买家昵评","评论内容","SKU/口味","评论日期","图片数"]
        widths = [6, 20, 25, 18, 15, 60, 15, 14, 8]
        data = all_reviews
    else:
        safe = re.sub(r'[^\w\u4e00-\u9fff]', '_', name).strip('_')[:20]
        filepath = os.path.join(OUTPUT_DIR, f"reviews_{safe}_{ts}.xlsx")
        ws_title = "评论数据"
        headers = ["序号","商品名称","商品ID","买家昵称","评论内容","SKU/口味","评论日期","图片数"]
        widths = [6, 25, 18, 15, 60, 15, 14, 8]
        data = reviews
    
    wb = Workbook()
    ws = wb.active
    ws.title = ws_title
    
    hf = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    hfl = PatternFill(start_color="FF2F5496", end_color="FF2F5496", fill_type="solid")
    ha = Alignment(horizontal="center", vertical="center")
    df = Font(name="Arial", size=10)
    da = Alignment(vertical="top", wrap_text=True)
    bd = Border(left=Side("thin","FFD0D0D0"), right=Side("thin","FFD0D0D0"),
                top=Side("thin","FFD0D0D0"), bottom=Side("thin","FFD0D0D0"))
    
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font, c.fill, c.alignment, c.border = hf, hfl, ha, bd
        ws.column_dimensions[get_column_letter(col)].width = w
    
    for i, r_val in enumerate(data, 1):
        if is_merged:
            row = [i, r_val.get("source_name",""), r_val.get("product_name",""),
                   r_val.get("product_id",""), r_val.get("nick",""),
                   r_val.get("content",""), r_val.get("sku",""),
                   r_val.get("date",""), r_val.get("pics","0")]
        else:
            row = [i, r_val.get("product_name",""), r_val.get("product_id",""),
                   r_val.get("nick",""), r_val.get("content",""),
                   r_val.get("sku",""), r_val.get("date",""), r_val.get("pics","0")]
        for col, val in enumerate(row, 1):
            c = ws.cell(row=i+1, column=col, value=val)
            c.font, c.border = df, bd
            c.alignment = da if (col == 6 if is_merged else col == 5) else ha
    
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(data)+1}"
    wb.save(filepath)
    return filepath


def main():
    print("=" * 70)
    print("  帝泊洱评论爬虫 v5 — Drawer评论区适配")
    print(f"  共 {len(PRODUCTS)} 个商品, 每个最多 {MAX_PER_PRODUCT} 条")
    print("=" * 70)
    
    all_reviews, file_list = [], []
    
    with sync_playwright() as p:
        context = create_browser(p)
        page = context.new_page()
        
        try:
            for idx, (name, url) in enumerate(PRODUCTS, 1):
                print(f"\n{'='*70}")
                print(f"  [{idx}/{len(PRODUCTS)}] {name}")
                print(f"{'='*70}")
                
                reviews = scrape_product(page, name, url)
                pid = re.search(r'id=(\d+)', page.url)
                pid = pid.group(1) if pid else f"prod_{idx}"
                
                for r in reviews:
                    r["product_name"] = name
                    r["product_id"] = pid
                    r["source_name"] = name
                
                if reviews:
                    fp = save_excel(reviews, name, pid)
                    file_list.append((name, fp, len(reviews)))
                    log(f"  💾 {os.path.basename(fp)}")
                
                all_reviews.extend(reviews)
        
        except KeyboardInterrupt:
            print("\n\n⚠️ 用户中断")
        except Exception as e:
            print(f"\n⚠️ 错误: {e}")
            import traceback; traceback.print_exc()
        finally:
            context.close()
            print("\n🔒 浏览器已关闭")
    
    # 总去重
    seen, unique = set(), []
    for r in all_reviews:
        k = (r.get("product_id",""), r.get("nick",""), r.get("content","")[:50])
        if k not in seen:
            seen.add(k)
            unique.append(r)
    
    if unique:
        merged = save_excel(None, None, None, is_merged=True, all_reviews=unique)
        print(f"\n{'='*70}")
        print(f"  📊 总计: {len(unique)} 条评论")
        for nm, fp, cnt in file_list:
            sz = os.path.getsize(fp) / 1024
            print(f"     📄 {nm}: {cnt}条 ({sz:.0f}KB)")
        print(f"  📁 合并: {os.path.basename(merged)} ({os.path.getsize(merged)/1024:.0f}KB)")
        print(f"  📂 位置: {OUTPUT_DIR}")
        print(f"{'='*70}")
    else:
        print("\n❌ 未获取到任何评论")

if __name__ == "__main__":
    main()
