#!/usr/bin/env python3
"""
天猫评论爬虫 v2 — 支持店铺全商品爬取

用法:
  python3 run.py                                    # 默认商品链接
  python3 run.py <商品链接>                          # 爬单个商品
  python3 run.py --shop <店铺链接>                   # 爬店铺所有商品
  python3 run.py --max 200                           # 限制每个商品最多200条
  python3 run.py --list-products <店铺链接>           # 仅列出商品链接，不爬评论
  python3 run.py --login-only                        # 仅打开浏览器登录天猫
"""

import sys, os, re, time, json, logging

logger = logging.getLogger(__name__)

# 接入统一的浏览器配置和登录管理
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from common.browser import STEALTH_JS, DEFAULT_USER_AGENT, DEFAULT_VIEWPORT, BROWSER_ARGS
from common.login_helper import (
    ensure_session, get_smart_headless, is_session_valid,
    mark_session_valid, get_profile_dir,
)

try:
    from playwright.sync_api import sync_playwright, TimeoutError as PwTimeout
except ImportError:
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "playwright", "-q"], check=True)
    from playwright.sync_api import sync_playwright, TimeoutError as PwTimeout

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl", "-q"], check=True)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "output")
PROFILE_DIR = str(get_profile_dir("tmall"))  # 统一 profile 目录
os.makedirs(OUTPUT_DIR, exist_ok=True)

# 天猫平台标识（供 login_helper 使用）
PLATFORM = "tmall"

# ===== 配置 =====
DEFAULT_URL = "https://detail.tmall.com/item.htm?id=686577281320"
MAX_PRODUCTS = 30
MAX_COMMENTS = 300
SCROLL_ROUNDS = 50
LOGIN_TIMEOUT = 180  # 3分钟等登录
TAB_CLICK_TIMEOUT = 10  # 每个tab点开后等几秒

SHOP_URL = "https://deepure.tmall.com/shop/view_shop.htm"


# ===== 工具 =====

def log(msg):
    print(f"  {msg}")

def extract_id(url):
    m = re.search(r"[?&]id=(\d+)", url)
    return m.group(1) if m else "unknown"

def extract_name(url):
    m = re.search(r"id=(\d+)", url)
    return f"商品_{m.group(1)}" if m else url.split("/")[-1][:30]

def parse_args():
    """解析命令行参数"""
    args = sys.argv[1:]
    mode = "single"
    target_url = DEFAULT_URL
    max_comments = MAX_COMMENTS
    list_only = False
    headless = False
    output_dir = OUTPUT_DIR
    login_only = False

    i = 0
    while i < len(args):
        if args[i] == "--shop":
            mode = "shop"
            if i + 1 < len(args) and not args[i + 1].startswith("--"):
                i += 1
                target_url = args[i]
        elif args[i] == "--max":
            i += 1
            max_comments = int(args[i])
        elif args[i] == "--list-products":
            list_only = True
            mode = "shop"
            if i + 1 < len(args) and not args[i + 1].startswith("--"):
                i += 1
                target_url = args[i]
        elif args[i] == "--headless":
            headless = True
        elif args[i] == "--login-only":
            login_only = True
        elif args[i] == "--output":
            i += 1
            output_dir = args[i]
        elif args[i].startswith("http"):
            target_url = args[i]
        i += 1

    return mode, target_url, max_comments, list_only, headless, output_dir, login_only


# ===== 浏览器 =====

def create_context(p, headless=False):
    """创建浏览器上下文（使用统一的反检测配置）"""
    return p.chromium.launch_persistent_context(
        PROFILE_DIR, headless=headless,
        args=BROWSER_ARGS,  # 使用统一的反检测参数
        viewport=DEFAULT_VIEWPORT,
        user_agent=DEFAULT_USER_AGENT,
        locale="zh-CN",
        bypass_csp=True,
    )


def wait_for_login(page, url, label="天猫"):
    """打开页面，等待用户手动登录（使用统一 session 管理）"""
    log(f"📄 正在打开 {label}...")
    page.goto(url, timeout=30000, wait_until="domcontentloaded")
    page.wait_for_timeout(3000)

    # 使用统一的登录检测
    return ensure_session(page, PLATFORM)


# ===== JS 提取店铺商品链接 =====

def extract_shop_products_js(page):
    """从页面上提取所有商品链接（浏览器端 JS 执行）"""
    js = """
    () => {
      const results = [];
      const seen = new Set();
      const links = document.querySelectorAll('a[href*="item.htm?id="]');
      links.forEach(a => {
        const href = a.href;
        const idMatch = href.match(/id=(\\d+)/);
        if (idMatch && !seen.has(idMatch[1])) {
          seen.add(idMatch[1]);
          const name = a.getAttribute('title')
            || (a.querySelector('img')?.getAttribute('alt'))
            || '';
          results.push({
            url: href.split('?')[0] + '?id=' + idMatch[1],
            name: name.trim().slice(0, 60),
            id: idMatch[1]
          });
        }
      });
      return results;
    }
    """
    return page.evaluate(js)


def scrape_shop_products(page, shop_url):
    """从店铺页提取所有商品链接，尝试切换分类tab"""
    all_products = []
    seen_ids = set()

    if not wait_for_login(page, shop_url, "帝泊洱店铺"):
        return []

    # 先试一次直接提取
    products = extract_shop_products_js(page)
    for p in products:
        if p["id"] not in seen_ids:
            seen_ids.add(p["id"])
            all_products.append(p)

    if all_products:
        log(f"   ✓ 首页找到 {len(all_products)} 个商品")

    # 尝试点击导航tab获取更多商品
    tab_selectors = [
        "text=自饮装", "text=分享装", "text=量贩装", "text=经典礼盒",
        "text=全部宝贝", "text=所有宝贝", "text=全部商品",
        "[class*='J_Tab'] a", "[class*='category'] a",
        "[class*='nav'] a", "[class*='menu'] a",
    ]

    for sel in tab_selectors:
        try:
            tabs = page.locator(sel)
            count = tabs.count()
            if count > 0:
                for idx in range(count):
                    try:
                        tab_text = tabs.nth(idx).inner_text().strip()
                        if tab_text in ("首页", "店铺首页", ""):
                            continue
                        log(f"   🔄 切换分类: {tab_text}")
                        tabs.nth(idx).click()
                        page.wait_for_timeout(TAB_CLICK_TIMEOUT * 1000)

                        products = extract_shop_products_js(page)
                        new_count = 0
                        for p in products:
                            if p["id"] not in seen_ids:
                                seen_ids.add(p["id"])
                                all_products.append(p)
                                new_count += 1
                        if new_count > 0:
                            log(f"      +{new_count} 个新商品 (累计 {len(all_products)})")
                    except Exception as e:
                        logger.warning(f"切换分类失败: {e}")
        except Exception as e:
            logger.warning(f"遍历分类失败: {e}")

    log(f"   ✅ 共 {len(all_products)} 个商品")
    return all_products


# ===== 评论提取（浏览器端） =====

def extract_reviews_js(page):
    js = """
    () => {
      const selectors = [
        "[class*='Comment--']",
        "[class*='commentItem']",
        "[class*='rate-'] [class*='item']",
        ".rate-content .rate-item",
        "[class*='review'] [class*='item']",
      ];
      let containers = [];
      for (const sel of selectors) {
        const els = document.querySelectorAll(sel);
        if (els.length > 0) { containers = els; break; }
      }
      if (containers.length === 0) {
        containers = Array.from(document.querySelectorAll('*')).filter(el => {
          const t = el.innerText || '';
          return t.includes('评论内容') && t.length > 50 && el.children.length < 10;
        });
      }
      return Array.from(containers).map(el => {
        const raw = el.innerText || '';
        const dateMatch = raw.match(/(\\d{4}[-\\.\\/]\\d{2}[-\\.\\/]\\d{2})/);
        const skuMatch = raw.match(/已购[：:]?\\s*(.+?)(?:\\n|$)/);
        const imgEls = el.querySelectorAll('img');
        const lines = raw.split('\\n').filter(l => l.trim());
        const nick = lines.length > 0 ? lines[0].trim().slice(0, 20) : '';
        let content = raw;
        const replyIdx = content.indexOf('商家回复');
        if (replyIdx > 0) content = content.slice(0, replyIdx);
        if (nick) content = content.replace(nick, '');
        content = content.replace(/\\d{4}[-\\.\\/]\\d{2}[-\\.\\/]\\d{2}.*?(?:\\n|$)/, '');
        content = content.replace(/已购[：:]?.*?(?:\\n|$)/, '');
        content = content.replace(/\\n\\s*\\n/g, '\\n').trim();
        if (!content) {
          content = lines.filter(l => l.length > 5 && !l.includes('已购') && !dateMatch)?.[0] || '';
        }
        return {
          date: dateMatch ? dateMatch[1] : '',
          nick: nick,
          content: content.slice(0, 500),
          sku: skuMatch ? skuMatch[1].trim() : '',
          pics: String(imgEls.length),
        };
      }).filter(r => r.content && r.content.length > 2);
    }
    """
    return page.evaluate(js)


# ===== 爬单个商品 =====

def scrape_product_reviews(page, url, product_name=None, max_comments=MAX_COMMENTS):
    item_id = extract_id(url)
    name = product_name or extract_name(url)
    log(f"\n📦 [{name}] 打开商品页...")

    page.goto(url, timeout=30000, wait_until="domcontentloaded")
    page.wait_for_timeout(3000)

    # 评价标签
    for sel in ["text=评价", "text=用户评价", "text=宝贝评价",
                 "[class*='tabDetailItem']:has-text('评价')"]:
        try:
            el = page.locator(sel).first
            if el.is_visible(timeout=3000):
                el.click()
                page.wait_for_timeout(2000)
                break
        except Exception as e:
            logger.warning(f"点击评价标签失败: {e}")

    # 关闭弹窗
    for sel in ["[class*='J_MIDDLEWARE']", "[class*='overlay'] .close",
                 "[class*='dialog'] .close", "[class*='modal'] .close"]:
        try:
            el = page.locator(sel).first
            if el.is_visible(timeout=800):
                el.click()
        except Exception as e:
            logger.warning(f"关闭弹窗失败: {e}")

    # 展开全部评价
    for sel in ["[class*='ShowButton']", "text=查看全部评价",
                 "text=全部评价", "text=查看更多"]:
        try:
            btn = page.locator(sel).first
            if btn.is_visible(timeout=2000):
                btn.click()
                page.wait_for_timeout(2000)
                break
        except Exception as e:
            logger.warning(f"展开全部评价失败: {e}")

    # 按最新排序
    try:
        btn = page.locator("text=最新").first
        if btn.is_visible(timeout=1500):
            btn.click()
            page.wait_for_timeout(1500)
    except Exception as e:
        logger.warning(f"按最新排序失败: {e}")

    # 滚动加载
    log(f"   📜 滚动加载 (上限 {max_comments} 条)...")
    last_count = 0
    no_change = 0
    for rnd in range(SCROLL_ROUNDS):
        page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        page.wait_for_timeout(1000)
        page.evaluate("window.scrollTo(0, document.body.scrollHeight - 300)")
        page.wait_for_timeout(500)
        page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        page.wait_for_timeout(800)

        reviews = extract_reviews_js(page)
        cur = len(reviews)

        if cur == last_count:
            no_change += 1
            if no_change >= 5:
                log(f"   ✓ {cur} 条")
                break
        else:
            no_change = 0
            if rnd % 3 == 0:
                log(f"   ... {cur} 条")
        last_count = cur
        if cur >= max_comments:
            log(f"   ✓ 已达上限 {max_comments} 条")
            break

    # 去重
    reviews = extract_reviews_js(page)
    seen = set()
    unique = []
    for r in reviews:
        k = (r.get("nick",""), r.get("content","")[:40])
        if r.get("content","") and k not in seen:
            seen.add(k)
            unique.append(r)

    log(f"   ✅ {len(unique)} 条有效评论")
    return unique


# ===== 保存 Excel =====

def save_to_excel(all_reviews, filepath):
    wb = Workbook()
    ws = wb.active
    ws.title = "评论数据"

    hf = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    hfl = PatternFill(start_color="FF2F5496", end_color="FF2F5496", fill_type="solid")
    ha = Alignment(horizontal="center", vertical="center")
    df = Font(name="Arial", size=10)
    da = Alignment(vertical="top", wrap_text=True)
    bd = Border(left=Side("thin","FFD0D0D0"), right=Side("thin","FFD0D0D0"),
                top=Side("thin","FFD0D0D0"), bottom=Side("thin","FFD0D0D0"))

    headers = ["序号","商品名称","商品ID","买家昵称","评论内容","SKU/口味","评论日期","图片数"]
    widths = [6, 30, 18, 15, 60, 15, 14, 8]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font, c.fill, c.alignment, c.border = hf, hfl, ha, bd
        ws.column_dimensions[chr(64+col)].width = w

    for i, r in enumerate(all_reviews, 1):
        row_data = [
            i,
            r.get("product_name",""), r.get("product_id",""),
            r.get("nick",""), r.get("content",""),
            r.get("sku",""), r.get("date",""), r.get("pics","0"),
        ]
        for col, val in enumerate(row_data, 1):
            c = ws.cell(row=i+1, column=col, value=val)
            c.font, c.border = df, bd
            c.alignment = da if col == 5 else ha

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{len(all_reviews)+1}"
    wb.save(filepath)
    return filepath


# ===== 主程序 =====

def main():
    mode, target_url, max_comments, list_only, headless, output_dir, login_only = parse_args()

    # 智能决定是否 headless
    headless = get_smart_headless(PLATFORM, headless)

    # login-only 模式
    if login_only:
        print("\n  🔑 登录模式：仅打开浏览器登录天猫\n")
        with sync_playwright() as p:
            context = create_context(p, headless=False)
            page = context.new_page()
            page.add_init_script(STEALTH_JS)
            try:
                ensure_session(page, PLATFORM)
                print("  ✅ 登录完成，session 已缓存")
            finally:
                context.close()
        return

    mode_str = {"shop": "店铺全商品", "single": "单商品"}[mode]
    print(f"\n{'='*70}")
    print(f"  天猫评论爬虫 v2  —  {mode_str}")
    print(f"  目标: {target_url}")
    print(f"  每个商品最多 {max_comments} 条评论")
    print(f"{'='*70}")

    os.makedirs(output_dir, exist_ok=True)

    # pre-check browser
    import subprocess
    r = subprocess.run([sys.executable, "-m", "playwright", "install", "--dry-run", "chromium"],
                        capture_output=True, text=True)
    if "already" not in r.stdout and "already" not in r.stderr:
        print("\n⚠️ 安装 Playwright Chromium...")
        subprocess.run([sys.executable, "-m", "playwright", "install", "chromium"], check=True)
        print("✅ 完成")

    all_reviews = []

    with sync_playwright() as p:
        context = create_context(p, headless=headless)
        page = context.new_page()
        page.add_init_script(STEALTH_JS)  # 注入反检测 JS

        try:
            if mode == "shop":
                products = scrape_shop_products(page, target_url)
                products = products[:MAX_PRODUCTS]

                if not products:
                    print("\n⚠️ 未找到商品")
                    return

                if list_only:
                    print(f"\n{'='*70}")
                    print(f"  店铺商品列表 ({len(products)} 个):")
                    for idx, prod in enumerate(products, 1):
                        print(f"  {idx}. {prod['name']}")
                        print(f"     {prod['url']}")
                    print(f"{'='*70}")
                    return

                print(f"\n{'='*70}")
                print(f"  开始爬取 {len(products)} 个商品")
                print(f"{'='*70}")

                for idx, prod in enumerate(products, 1):
                    print(f"\n--- [{idx}/{len(products)}] {prod['name']} ---")
                    reviews = scrape_product_reviews(page, prod["url"], prod["name"], max_comments)
                    for r in reviews:
                        r["product_name"] = prod["name"]
                        r["product_id"] = prod["id"]
                    all_reviews.extend(reviews)
                    print(f"   📊 累计: {len(all_reviews)} 条")

            else:
                reviews = scrape_product_reviews(page, target_url, max_comments=max_comments)
                pid = extract_id(target_url)
                for r in reviews:
                    r["product_name"] = extract_name(target_url)
                    r["product_id"] = pid
                all_reviews = reviews

        except Exception as e:
            print(f"\n⚠️ 错误: {e}")
            import traceback; traceback.print_exc()
        finally:
            context.close()
            print("\n🔒 浏览器已关闭")

    # 总去重
    seen = set()
    unique = []
    for r in all_reviews:
        k = (r.get("product_id",""), r.get("nick",""), r.get("content","")[:50])
        if k not in seen:
            seen.add(k)
            unique.append(r)

    # 保存
    ts = time.strftime("%Y%m%d_%H%M%S")
    fp = os.path.join(output_dir, f"reviews_{mode}_{ts}.xlsx")
    save_to_excel(unique, fp)

    # 统计
    stats = {}
    for r in unique:
        pn = r.get("product_name","未知")
        stats[pn] = stats.get(pn, 0) + 1

    print(f"\n{'='*70}")
    print(f"  ✅ 爬取完成")
    print(f"  共 {len(unique)} 条评论")
    for name, count in stats.items():
        print(f"    {name}: {count} 条")
    print(f"\n  📊 {fp}")
    print(f"{'='*70}\n")


if __name__ == "__main__":
    main()
