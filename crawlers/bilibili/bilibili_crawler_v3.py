#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
B站视频数据爬虫 v3.0
- 正常表格格式输出
- 支持动态数量的共创参与者（最多10个UP主）
- 支持时间对比（每次爬取记录时间戳）
- 一键启动：python3 bilibili_crawler_v3.py

用法:
  python3 bilibili_crawler_v3.py BV1Ww411175v BV号2 ...
  python3 bilibili_crawler_v3.py --file urls.txt
"""

import ssl
import urllib.request
import json
import sys
import os
import re
import logging
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# 接入统一的 UA 配置
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from common.browser import DEFAULT_USER_AGENT

# ============ 配置 ============
MAX_UP_COUNT = 10  # 最多支持10个共创参与者

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE

HEADERS = {
    "User-Agent": DEFAULT_USER_AGENT,
    "Referer": "https://www.bilibili.com",
    "Accept": "application/json",
    "Cookie": "buvid3=bilibili_crawler_v3; b_nut=1700000000",
}

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s", datefmt="%H:%M:%S")
logger = logging.getLogger("bilibili_crawler_v3")

# 输出文件（项目 output 目录，不再放桌面）
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "bilibili_data.xlsx")


def build_headers():
    """动态生成表头"""
    headers = [
        "抓取时间",
        "BV号",
        "AID",
        "视频标题",
        "发布时间",
        "时长(秒)",
        "时长(分:秒)",
        "分区",
        "播放量",
        "点赞数",
        "投币数",
        "收藏数",
        "分享数",
        "弹幕数",
        "评论数",
        "参与人数",
    ]
    # 动态添加UP主列
    for i in range(1, MAX_UP_COUNT + 1):
        headers.extend([
            f"UP{i}_UID",
            f"UP{i}_昵称",
            f"UP{i}_角色",
            f"UP{i}_粉丝数",
        ])
    return headers


HEADERS_DEF = build_headers()


# ============ API 函数 ============

def http_get(url: str, retries: int = 3) -> dict:
    """带SSL的HTTP GET请求，带重试机制"""
    import time
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers=HEADERS)
            with urllib.request.urlopen(req, timeout=15, context=ctx) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except Exception as e:
            if attempt < retries - 1:
                print(f"    ⏳ 重试 ({attempt + 1}/{retries})...")
                time.sleep(2)
            else:
                raise e


def get_video_info(bvid):
    """获取视频详细信息"""
    url = f'https://api.bilibili.com/x/web-interface/view?bvid={bvid}'
    try:
        data = http_get(url)
        if data['code'] != 0:
            print(f"  ❌ API错误: {data.get('message', '未知')}")
            return None
        return data['data']
    except Exception as e:
        print(f"  ❌ 网络错误: {e}")
        return None


def get_up_fans(mid):
    """获取UP主粉丝数"""
    url = f'https://api.bilibili.com/x/relation/stat?vmid={mid}'
    try:
        data = http_get(url)
        if data.get("code") == 0:
            return data['data']['follower']
    except Exception as e:
        logger.warning(f"获取粉丝数失败: {e}")
    return 0


def parse_duration(seconds):
    """秒数转分:秒格式"""
    m, s = divmod(seconds, 60)
    return f"{m}:{s:02d}"


def extract_video_data(video_info):
    """从视频信息提取数据"""
    v = video_info
    stat = v.get('stat', {})
    owner = v.get('owner', {})
    staff = v.get('staff', []) or []

    # 获取所有参与者
    all_participants = []

    # 主UP主作为第一个参与者
    if owner.get('mid'):
        all_participants.append({
            'uid': owner.get('mid'),
            '昵称': owner.get('name', ''),
            '角色': 'UP主',
        })

    # 共创参与者（排除主UP主）
    other_staff = [s for s in staff if str(s.get('mid', '')) != str(owner.get('mid', ''))]
    for s in other_staff:
        all_participants.append({
            'uid': s.get('mid', 0),
            '昵称': s.get('name', ''),
            '角色': s.get('title', ''),
        })

    # 填充到最大数量（用空值）
    while len(all_participants) < MAX_UP_COUNT:
        all_participants.append({'uid': '', '昵称': '', '角色': '', '粉丝数': ''})

    # 构建行数据
    row_data = {
        "抓取时间": datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        "BV号": v.get('bvid', ''),
        "AID": v.get('aid', 0),
        "视频标题": v.get('title', ''),
        "发布时间": datetime.fromtimestamp(v.get('pubdate', 0)).strftime('%Y-%m-%d %H:%M'),
        "时长(秒)": v.get('duration', 0),
        "时长(分:秒)": parse_duration(v.get('duration', 0)),
        "分区": v.get('tname', ''),
        "播放量": stat.get('view', 0),
        "点赞数": stat.get('like', 0),
        "投币数": stat.get('coin', 0),
        "收藏数": stat.get('favorite', 0),
        "分享数": stat.get('share', 0),
        "弹幕数": stat.get('danmaku', 0),
        "评论数": stat.get('reply', 0),
        "参与人数": len(staff) if staff else 1,
    }

    # 填充UP主信息
    for i, p in enumerate(all_participants[:MAX_UP_COUNT], 1):
        row_data[f"UP{i}_UID"] = p.get('uid', '')
        row_data[f"UP{i}_昵称"] = p.get('昵称', '')
        row_data[f"UP{i}_角色"] = p.get('角色', '')
        row_data[f"UP{i}_粉丝数"] = p.get('粉丝数', '')

    # 获取粉丝数
    for i, p in enumerate(all_participants[:MAX_UP_COUNT], 1):
        if p.get('uid'):
            row_data[f"UP{i}_粉丝数"] = get_up_fans(p['uid'])

    return row_data


# ============ Excel 操作 ============

def init_excel(output_file=OUTPUT_FILE):
    """初始化Excel文件（如果不存在）"""
    os.makedirs(os.path.dirname(output_file), exist_ok=True)

    if not os.path.exists(output_file):
        wb = Workbook()
        ws = wb.active
        ws.title = "数据汇总"

        # 写入表头
        ws.append(HEADERS_DEF)

        # 表头样式
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=10)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col, header in enumerate(HEADERS_DEF, 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align

        # 设置列宽
        col_widths = [18, 14, 11, 45, 16, 8, 8, 10, 12, 10, 10, 10, 10, 10, 10, 8]  # 基础列
        for i in range(MAX_UP_COUNT):  # UP主列
            col_widths.extend([12, 14, 8, 12])

        for i, width in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        # 冻结首行
        ws.freeze_panes = "A2"

        wb.save(output_file)
        print(f"📁 新建文件: {output_file}")

    return load_workbook(output_file)


def append_to_excel(data_list, output_file=OUTPUT_FILE):
    """追加数据到Excel"""
    wb = init_excel(output_file)
    ws = wb.active

    for data in data_list:
        row = [data.get(h, "") for h in HEADERS_DEF]
        ws.append(row)

        # 数据行样式
        last_row = ws.max_row
        for col in range(1, len(HEADERS_DEF) + 1):
            cell = ws.cell(row=last_row, column=col)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    wb.save(output_file)
    wb.close()


# ============ 解析 BV 号 ============

def parse_bvid(arg):
    """从URL或纯文本解析BV号（保持原始大小写）"""
    if 'BV' in arg.upper():
        match = re.search(r'BV[a-zA-Z0-9]+', arg)
        if match:
            return match.group()
    if re.match(r'^[a-zA-Z0-9]+$', arg):
        return f'BV{arg}'
    return None


def parse_urls_file(filepath):
    """从文件解析BV号列表"""
    bvids = []
    with open(filepath, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith('#'):
                bvid = parse_bvid(line)
                if bvid:
                    bvids.append(bvid)
    return bvids


# ============ 主程序 ============

def main():
    print("\n" + "="*50)
    print("📺 B站视频数据爬虫 v3.0")
    print("="*50 + "\n")

    if len(sys.argv) < 2:
        print("用法:")
        print("  python3 bilibili_crawler_v3.py BV号1 BV号2 ...")
        print("  python3 bilibili_crawler_v3.py --file urls.txt")
        print("\n示例:")
        print("  python3 bilibili_crawler_v3.py BV1Ww411175v")
        print("  python3 bilibili_crawler_v3.py https://www.bilibili.com/video/BV1Ww411175v")
        sys.exit(1)

    # 解析参数
    bvids = []
    output_file = OUTPUT_FILE
    args = sys.argv[1:]
    i = 0
    while i < len(args):
        if args[i] == '--file':
            if i + 1 < len(args):
                bvids = parse_urls_file(args[i + 1])
            i += 2
            continue
        elif args[i] == '--output':
            if i + 1 < len(args):
                output_file = args[i + 1]
            i += 2
            continue
        else:
            bvid = parse_bvid(args[i])
            if bvid:
                bvids.append(bvid)
            i += 1

    if not bvids:
        print("❌ 未找到有效的BV号")
        sys.exit(1)

    # 去重
    bvids = list(dict.fromkeys(bvids))

    print(f"📋 待爬取: {len(bvids)} 个视频\n")

    all_data = []

    for i, bvid in enumerate(bvids, 1):
        print(f"[{i}/{len(bvids)}] 🔍 {bvid}")

        video_info = get_video_info(bvid)
        if not video_info:
            continue

        data = extract_video_data(video_info)

        # 显示参与者摘要
        participant_count = data['参与人数']
        print(f"    → 共 {participant_count} 位参与者")

        # 输出前3个参与者
        for j in range(1, min(4, participant_count + 1)):
            name = data.get(f'UP{j}_昵称', '')
            role = data.get(f'UP{j}_角色', '')
            fans = data.get(f'UP{j}_粉丝数', 0)
            if name:
                print(f"       {j}. {name} ({role}) - {fans:,}粉丝")

        # 输出摘要
        print(f"    ✅ {data['视频标题'][:30]}...")
        print(f"       👁 {data['播放量']:,}  👍 {data['点赞数']:,}  ⭐ {data['收藏数']:,}")

        all_data.append(data)

    if all_data:
        print(f"\n💾 保存到: {output_file}")
        append_to_excel(all_data, output_file)
        print(f"\n🎉 完成！共爬取 {len(all_data)} 条数据")
        print(f"   数据已追加到 Excel，现在可以打开查看\n")
    else:
        print("\n❌ 未获取到任何数据")


if __name__ == '__main__':
    main()
