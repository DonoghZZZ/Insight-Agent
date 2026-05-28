#!/usr/bin/env python3
"""
抖音视频爬虫 v2 — 抓取视频内容 + 评论区
用法:
  python3 douyin_scraper.py                                    # 默认视频URL
  python3 douyin_scraper.py <抖音视频URL>                      # 指定视频URL
  python3 douyin_scraper.py --max 50                           # 最多50条评论
  python3 douyin_scraper.py --local <本地HTML文件路径>          # 解析本地HTML
  python3 douyin_scraper.py --urls urls.txt                    # 批量爬取
  python3 douyin_scraper.py --headless                         # 无头模式
"""

import sys
import os
import re
import time
import json
import logging
import argparse
from typing import List, Dict, Any, Optional
from datetime import datetime
import threading
import random

# ========== 终端颜色和样式定义 ==========
class TerminalColors:
    RESET = "\033[0m"
    BOLD = "\033[1m"
    DIM = "\033[2m"
    ITALIC = "\033[3m"
    UNDERLINE = "\033[4m"
    BLINK = "\033[5m"
    REVERSE = "\033[7m"
    HIDDEN = "\033[8m"

    # 前景色
    BLACK = "\033[30m"
    RED = "\033[31m"
    GREEN = "\033[32m"
    YELLOW = "\033[33m"
    BLUE = "\033[34m"
    MAGENTA = "\033[35m"
    CYAN = "\033[36m"
    WHITE = "\033[37m"
    BRIGHT_BLACK = "\033[90m"
    BRIGHT_RED = "\033[91m"
    BRIGHT_GREEN = "\033[92m"
    BRIGHT_YELLOW = "\033[93m"
    BRIGHT_BLUE = "\033[94m"
    BRIGHT_MAGENTA = "\033[95m"
    BRIGHT_CYAN = "\033[96m"
    BRIGHT_WHITE = "\033[97m"

    # 背景色
    BG_BLACK = "\033[40m"
    BG_RED = "\033[41m"
    BG_GREEN = "\033[42m"
    BG_YELLOW = "\033[43m"
    BG_BLUE = "\033[44m"
    BG_MAGENTA = "\033[45m"
    BG_CYAN = "\033[46m"
    BG_WHITE = "\033[47m"

TC = TerminalColors

# ========== 炫酷装饰文案 ==========
COOL_PHRASES = [
    "🚀 正在向数据宇宙进发...",
    "⚡ 代码正在高速运行中...",
    "🎯 精准锁定目标内容...",
    "🔥 数据正在熊熊燃烧...",
    "✨ 魔法正在施展中...",
    "🌟 精彩内容即将呈现...",
    "💎 珍贵数据正在挖掘...",
    "🎪 欢迎来到数据马戏团...",
    "🎭 爬虫正在卖力表演...",
    "🎡 数据摩天轮旋转中...",
]

SPINNER_CHARS = ["⠋", "⠙", "⠹", "⠸", "⠼", "⠴", "⠦", "⠧", "⠇", "⠏"]
PROGRESS_BARS = ["▏", "▎", "▍", "▌", "▋", "▊", "▉", "█"]

# ========== 炫酷进度显示类 ==========
class CoolProgress:
    def __init__(self):
        self.spinner_active = False
        self.spinner_thread = None
        self.current_spinner_idx = 0

    def print_banner(self):
        banner = f"""
{TC.BRIGHT_CYAN}{TC.BOLD}╔════════════════════════════════════════════════════════════════╗
{TC.BRIGHT_CYAN}{TC.BOLD}║{TC.RESET}  {TC.BRIGHT_MAGENTA}{TC.BOLD}🎬  抖 音 视 频 爬 虫  v2  🎬{TC.RESET}{' ' * 36}{TC.BRIGHT_CYAN}{TC.BOLD}║
{TC.BRIGHT_CYAN}{TC.BOLD}║{TC.RESET}  {TC.BRIGHT_BLUE}━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━{TC.RESET}  {TC.BRIGHT_CYAN}{TC.BOLD}║
{TC.BRIGHT_CYAN}{TC.BOLD}║{TC.RESET}  {TC.BRIGHT_GREEN}📹 视频信息    💬 评论数据    📊 Excel导出    ⚡ 高性能{TC.RESET}     {TC.BRIGHT_CYAN}{TC.BOLD}║
{TC.BRIGHT_CYAN}{TC.BOLD}╚════════════════════════════════════════════════════════════════╝{TC.RESET}
"""
        print(banner)

    def cool_print(self, message, style="info", end="\n"):
        styles = {
            "info": f"{TC.BRIGHT_CYAN}ℹ️ {TC.RESET}",
            "success": f"{TC.BRIGHT_GREEN}✅ {TC.RESET}",
            "warning": f"{TC.BRIGHT_YELLOW}⚠️ {TC.RESET}",
            "error": f"{TC.BRIGHT_RED}❌ {TC.RESET}",
            "rocket": f"{TC.BRIGHT_MAGENTA}🚀 {TC.RESET}",
            "star": f"{TC.BRIGHT_YELLOW}✨ {TC.RESET}",
            "fire": f"{TC.BRIGHT_RED}🔥 {TC.RESET}",
        }
        prefix = styles.get(style, "")
        print(f"{prefix}{message}", end=end)

    def start_spinner(self, message):
        self.spinner_active = True
        self.spinner_message = message
        self.spinner_thread = threading.Thread(target=self._spinner_loop)
        self.spinner_thread.daemon = True
        self.spinner_thread.start()

    def stop_spinner(self, success_message=None):
        self.spinner_active = False
        if self.spinner_thread:
            self.spinner_thread.join(timeout=0.5)
        if success_message:
            print(f"\r{TC.BRIGHT_GREEN}✅ {success_message}{TC.RESET}")
        else:
            print()

    def _spinner_loop(self):
        while self.spinner_active:
            char = SPINNER_CHARS[self.current_spinner_idx % len(SPINNER_CHARS)]
            phrase = random.choice(COOL_PHRASES)
            print(f"\r{TC.BRIGHT_MAGENTA}{char}{TC.RESET} {self.spinner_message} {TC.DIM}({phrase}){TC.RESET}", end="", flush=True)
            self.current_spinner_idx += 1
            time.sleep(0.1)

    def print_progress_bar(self, current, total, prefix="进度", width=50):
        percentage = (current / total) if total > 0 else 0
        filled = int(width * percentage)
        bar = "█" * filled + "░" * (width - filled)
        print(f"\r{TC.BRIGHT_CYAN}{prefix}{TC.RESET} |{TC.BRIGHT_GREEN}{bar}{TC.RESET}| {TC.BRIGHT_YELLOW}{percentage*100:.1f}%{TC.RESET} ({current}/{total})", end="", flush=True)

    def print_scrolling_text(self, text, speed=0.03):
        for char in text:
            print(char, end="", flush=True)
            time.sleep(speed)
        print()

    def print_step(self, step_num, total_steps, title, description=""):
        box_width = 60
        print(f"\n{TC.BRIGHT_BLUE}{'═' * box_width}{TC.RESET}")
        print(f"{TC.BRIGHT_MAGENTA}{TC.BOLD}  【步骤 {step_num}/{total_steps}】{TC.RESET} {TC.BRIGHT_CYAN}{title}{TC.RESET}")
        if description:
            print(f"  {TC.DIM}{description}{TC.RESET}")
        print(f"{TC.BRIGHT_BLUE}{'─' * box_width}{TC.RESET}\n")

cool_progress = CoolProgress()

# ========== 自动安装依赖 ==========
try:
    from playwright.sync_api import sync_playwright, TimeoutError as PwTimeout
except ImportError:
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "playwright", "-q"], check=True)
    from playwright.sync_api import sync_playwright, TimeoutError as PwTimeout

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl", "-q"], check=True)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

try:
    from bs4 import BeautifulSoup
except ImportError:
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "beautifulsoup4", "-q"], check=True)
    from bs4 import BeautifulSoup

# ========== 路径与常量 ==========
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(SCRIPT_DIR, ".."))
from common.browser import STEALTH_JS, DEFAULT_USER_AGENT, DEFAULT_VIEWPORT, BROWSER_ARGS
from common.login_helper import (
    ensure_session, get_smart_headless, is_session_valid,
    mark_session_valid, DOUYIN_CHECK_LOGIN_JS,
)
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)

DEFAULT_URL = "https://www.douyin.com/video/7631962423612765609"
MAX_COMMENTS = 10000
SCROLL_ROUNDS = 200
SCROLL_WAIT_MS = 1500
NO_CHANGE_THRESHOLD = 15
RETRY_TIMES = 3

# ========== 日志配置 ==========
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger("douyin_scraper")

def log(msg: str, level: str = "info"):
    if level == "info":
        logger.info(msg)
    elif level == "warning":
        logger.warning(msg)
    elif level == "error":
        logger.error(msg)

# ========== 命令行参数 ==========
def parse_args():
    parser = argparse.ArgumentParser(
        description="抖音视频爬虫 — 抓取视频内容 + 评论区",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("url", nargs="?", default=DEFAULT_URL, help="抖音视频URL")
    parser.add_argument("--max", dest="max_comments", type=int, default=MAX_COMMENTS,
                        help=f"最多爬取评论数（默认: {MAX_COMMENTS}）")
    parser.add_argument("--local", dest="local_file", default=None, help="解析本地HTML文件")
    parser.add_argument("--urls", dest="urls", default=None, help="批量爬取：URL列表文件（每行一个）")
    parser.add_argument("--urls-file", dest="urls_file", default=None, help="批量爬取：URL列表文件（每行一个）")
    parser.add_argument("--output", dest="output_file", default=None, help="指定输出Excel文件路径")
    parser.add_argument("--headless", dest="headless", action="store_true", default=False,
                        help="无头模式运行（默认显示浏览器）")
    parser.add_argument("--login-only", dest="login_only", action="store_true", default=False,
                        help="仅打开浏览器登录，登录成功后退出（用于首次配置）")
    return parser.parse_args()

# ========== 浏览器配置（反检测） ==========
def create_browser_context(p, headless: bool = False):
    profile_dir = os.path.join(SCRIPT_DIR, "browser_profile_douyin")
    os.makedirs(profile_dir, exist_ok=True)
    context = p.chromium.launch_persistent_context(
        profile_dir,
        headless=headless,
        args=BROWSER_ARGS,
        viewport=DEFAULT_VIEWPORT,
        user_agent=DEFAULT_USER_AGENT,
        locale="zh-CN",
        timezone_id="Asia/Shanghai",
        bypass_csp=True,
    )
    return context

# ========== 反检测JS ==========
ANTI_DETECTION_JS = STEALTH_JS

# ========== 浏览器端JS：提取视频信息 ==========
EXTRACT_VIDEO_INFO_JS = """
() => {
  const result = {};
  const numRe = /([\\d.]+[万wW]?)/;

  // 1. 视频描述
  const h1 = document.querySelector('h1');
  if (h1 && h1.innerText.trim()) {
    result.desc = h1.innerText.trim();
  }

  // 2. 作者信息 - 优先用 data-e2e="user-info"
  const userInfoEl = document.querySelector('[data-e2e="user-info"]');
  if (userInfoEl) {
    const text = userInfoEl.innerText || '';
    const fanMatch = text.match(/粉丝\\s*([\\d.]+[万wW]?)/);
    const likeMatch = text.match(/获赞\\s*([\\d.]+[万wW]?)/);
    if (fanMatch) result.followers = fanMatch[1];
    if (likeMatch) result.totalLikes = likeMatch[1];
    
    const links = userInfoEl.querySelectorAll('a[href*="/user/"]');
    for (const link of links) {
      let linkText = link.innerText.trim().split('\\n')[0];
      linkText = linkText.replace(/\\s*作者$/, '').trim();
      linkText = linkText.replace(/认证徽章.*/, '').trim();
      linkText = linkText.replace(/优质视频创.*/, '').trim();
      linkText = linkText.replace(/视频创.*/, '').trim();
      if (linkText && linkText.length > 0 && linkText.length < 30 &&
          !linkText.includes('粉丝') && !linkText.includes('获赞') && linkText !== '作者') {
        result.author = linkText;
        break;
      }
    }
  }

  // 兜底：用头像img的alt
  if (!result.author) {
    const avatarImgs = document.querySelectorAll('img[alt*="头像"]');
    for (const img of avatarImgs) {
      if (img.closest('[data-e2e="comment-list"]')) continue;
      const alt = img.alt;
      const name = alt.replace(/的头像$/, '').replace(/头像$/, '').trim();
      if (name && name.length > 0 && name.length < 30) {
        result.author = name;
        break;
      }
    }
  }

  // 如果还没找到粉丝获赞，单独提取
  if (!result.followers || !result.totalLikes) {
    const allEls2 = document.querySelectorAll('*');
    for (const el of allEls2) {
      const text = el.innerText || '';
      if (text.includes('粉丝') && text.includes('获赞') && text.length < 50) {
        if (!result.followers) {
          const fanMatch = text.match(/粉丝\\s*([\\d.]+[万wW]?)/);
          if (fanMatch) result.followers = fanMatch[1];
        }
        if (!result.totalLikes) {
          const likeMatch = text.match(/获赞\\s*([\\d.]+[万wW]?)/);
          if (likeMatch) result.totalLikes = likeMatch[1];
        }
        if (result.followers && result.totalLikes) break;
      }
    }
  }

  // 3. 点赞/评论/收藏/转发 - 通过data-e2e属性提取
  const e2eMap = {
    'video-player-digg': 'likes',
    'video-player-comment': 'commentCount',
    'video-player-fav': 'favs',
    'video-player-share': 'shares'
  };
  for (const [e2e, key] of Object.entries(e2eMap)) {
    const el = document.querySelector(`[data-e2e="${e2e}"]`);
    if (el) {
      const text = el.innerText.trim();
      const m = text.match(numRe);
      if (m) result[key] = m[1];
    }
  }

  // 如果data-e2e没找到，用操作栏区域提取
  if (!result.likes) {
    const actionContainer = document.querySelector('[data-e2e="feed-comment-icon"]')?.parentElement?.parentElement;
    if (actionContainer) {
      const nums = [];
      for (const el of actionContainer.querySelectorAll(':scope > div')) {
        const text = el.innerText?.trim() || '';
        const m = text.match(numRe);
        if (m) nums.push(m[1]);
      }
      if (nums.length >= 1) result.likes = nums[0];
      if (nums.length >= 2) result.commentCount = nums[1];
      if (nums.length >= 3) result.favs = nums[2];
      if (nums.length >= 4) result.shares = nums[3];
    }
  }

  // 4. 发布时间
  const publishEl = document.querySelector('[data-e2e="detail-video-publish-time"]');
  if (publishEl) {
    result.publishTime = (publishEl.innerText || '').replace(/^发布时间[：:]\\s*/, '').trim();
  }

  return result;
}
"""

# ========== 浏览器端JS：提取评论列表 ==========
EXTRACT_COMMENTS_JS = """
() => {
    const results = [];
    const seen = new Set();
    const timeRe = /^(刚刚|昨天|\\d+年前|\\d+月前|\\d+天前|\\d+小时前|\\d+分钟前|\\d{4}-\\d{2}-\\d{2}|\\d+秒前)/;

    const extractSingleComment = (el, isChild) => {
        let name = '';
        let replyTo = '';
        let content = '';
        let time = '';
        let likes = 0;
        let replies = 0;
        let authorLiked = false;

        const infoWrap = el.querySelector('.comment-item-info-wrap');
        if (infoWrap) {
            const links = infoWrap.querySelectorAll('a[href*="/user/"]');
            if (isChild && links.length >= 2) {
                name = links[0].innerText.trim().split('\\n')[0].trim();
                replyTo = links[1].innerText.trim().split('\\n')[0].trim();
            } else if (links.length >= 1) {
                name = links[0].innerText.trim().split('\\n')[0].trim();
            }
            if (!name) {
                const firstLink = infoWrap.querySelector('a');
                if (firstLink) {
                    name = firstLink.innerText.trim().split('\\n')[0].trim();
                }
            }
            const tagEl = infoWrap.querySelector('.comment-item-tag');
            if (tagEl && /作者/.test(tagEl.innerText || '')) authorLiked = true;
        }

        const contentEl = el.querySelector('.gD6hDm2O');
        if (contentEl) {
            content = contentEl.innerText.trim();
        }
        if (!content) {
            const infoWrap2 = el.querySelector('.comment-item-info-wrap');
            if (infoWrap2 && infoWrap2.parentElement) {
                const parent = infoWrap2.parentElement;
                const children = [...parent.children];
                const infoIdx = children.indexOf(infoWrap2);
                const contentTexts = [];
                for (let i = infoIdx + 1; i < children.length; i++) {
                    const child = children[i];
                    const cls = typeof child.className === 'string' ? child.className : '';
                    if (cls.includes('comment-item-stats') || cls.includes('vo4kEeuY') || cls.includes('F3UVeHBA')) break;
                    const t = child.innerText?.trim() || '';
                    if (t && !timeRe.test(t) && t !== '分享' && t !== '回复') {
                        contentTexts.push(t);
                    }
                }
                content = contentTexts.join(' ').trim();
            }
        }

        const timeEl = el.querySelector('.vo4kEeuY');
        if (timeEl) time = timeEl.innerText.trim();
        if (!time) {
            const spans = el.querySelectorAll('span');
            for (const s of spans) {
                const t = s.innerText.trim();
                if (timeRe.test(t) && t.length < 20) { time = t; break; }
            }
        }

        const statsEl = el.querySelector('.comment-item-stats-container');
        if (statsEl) {
            const likeP = statsEl.querySelector('.soEq5p_Y');
            if (likeP) {
                const spans = likeP.querySelectorAll('span');
                for (const s of spans) {
                    const t = s.innerText.trim();
                    if (/^\\d+$/.test(t)) { likes = parseInt(t); break; }
                }
            }
        }

        const replyBtn = el.querySelector('.comment-reply-expand-btn');
        if (replyBtn) {
            const t = replyBtn.innerText.trim();
            const m = t.match(/展开(\\d+)条回复/);
            if (m) replies = parseInt(m[1]);
        }

        content = content.replace(/作者赞过/g, '').trim();
        if (replyTo) {
            content = '回复 ' + replyTo + ': ' + content;
        }

        const key = (name + '|' + content.slice(0, 80)).trim();
        if (key && !seen.has(key) && name) {
            seen.add(key);
            results.push({ name, content, likes, time, replies, authorLiked, isChild });
        }
    };

    const mainComments = document.querySelectorAll('[data-e2e="comment-item"]');
    mainComments.forEach(commentItem => {
        const mainContainer = commentItem.querySelector('.aKq6ctMW');
        if (!mainContainer) {
            extractSingleComment(commentItem, false);
            return;
        }

        const allHpZGj2UP = mainContainer.querySelectorAll(':scope > .hpZGj2UP');
        if (allHpZGj2UP.length === 0) {
            extractSingleComment(commentItem, false);
            return;
        }

        allHpZGj2UP.forEach((item, idx) => {
            const infoWrap = item.querySelector('.comment-item-info-wrap');
            if (!infoWrap) return;
            const hasReplyArrow = !!infoWrap.querySelector('.Ow9zILbw');
            const isChild = idx > 0 || hasReplyArrow;
            extractSingleComment(item, isChild);
        });
    });

    // 第二遍：查找独立的子评论项（展开后出现在评论列表中的子评论）
    // 这些子评论有独立的 data-e2e="comment-item"，但内部有回复箭头
    const allCommentItems = document.querySelectorAll('[data-e2e="comment-item"]');
    allCommentItems.forEach(commentItem => {
        const infoWrap = commentItem.querySelector('.comment-item-info-wrap');
        if (!infoWrap) return;
        
        // 检查是否有回复箭头（子评论标志）
        const hasReplyArrow = !!infoWrap.querySelector('.Ow9zILbw');
        if (!hasReplyArrow) return;
        
        // 检查是否已经被提取过
        const links = infoWrap.querySelectorAll('a[href*="/user/"]');
        let name = links.length >= 2 ? links[0].innerText.trim().split('\\n')[0].trim() : '';
        let replyTo = links.length >= 2 ? links[1].innerText.trim().split('\\n')[0].trim() : '';
        const contentEl = commentItem.querySelector('.gD6hDm2O');
        let content = contentEl ? contentEl.innerText.trim() : '';
        if (replyTo) content = '回复 ' + replyTo + ': ' + content;
        
        const key = (name + '|' + content.slice(0, 80)).trim();
        if (key && !seen.has(key) && name) {
            seen.add(key);
            const timeEl = commentItem.querySelector('.vo4kEeuY');
            const statsEl = commentItem.querySelector('.comment-item-stats-container');
            let likes = 0;
            if (statsEl) {
                const likeP = statsEl.querySelector('.soEq5p_Y');
                if (likeP) {
                    for (const s of likeP.querySelectorAll('span')) {
                        const t = s.innerText.trim();
                        if (/^\\d+$/.test(t)) { likes = parseInt(t); break; }
                    }
                }
            }
            const tagEl = infoWrap.querySelector('.comment-item-tag');
            results.push({
                name,
                content,
                likes,
                time: timeEl ? timeEl.innerText.trim() : '',
                replies: 0,
                authorLiked: tagEl ? /作者/.test(tagEl.innerText || '') : false,
                isChild: true
            });
        }
    });

    if (results.length === 0) {
        const fallbacks = document.querySelectorAll('[data-e2e="comment-item"]');
        fallbacks.forEach(el => {
            const fullText = (el.innerText || '').trim();
            if (fullText.length > 5) {
                const lines = fullText.split('\\n').filter(l => l.trim());
                const fn = lines[0] || '';
                const fc = lines.length > 1 ? lines.slice(1).join(' ') : fullText;
                const fk = (fn + '|' + fc.slice(0, 50)).trim();
                if (fk && !seen.has(fk)) {
                    seen.add(fk);
                    results.push({ name: fn, content: fc, likes: 0, time: '', replies: 0, authorLiked: false, isChild: false });
                }
            }
        });
    }

    return results;
}
"""

# ========== 本地HTML解析 ==========
def _extract_deep_text(el):
    """从多层嵌套的 span 中提取最深层纯文本"""
    if el is None:
        return ""
    spans = el.find_all("span", recursive=True)
    texts = []
    for s in spans:
        if s.find("img"):
            continue
        t = s.get_text(strip=True)
        if t and len(t) < 100:
            texts.append(t)
    if texts:
        return max(texts, key=len)
    return el.get_text(strip=True)


def parse_local_html(file_path: str, source_url: str = "") -> Dict[str, Any]:
    """解析本地保存的抖音HTML文件，提取视频信息 + 评论"""
    cool_progress.cool_print(f"解析本地HTML文件: {file_path}", "rocket")

    if not os.path.exists(file_path):
        raise FileNotFoundError(f"文件不存在: {file_path}")

    cool_progress.start_spinner("正在加载HTML文件...")
    with open(file_path, "r", encoding="utf-8") as f:
        html = f.read()
    cool_progress.stop_spinner("HTML文件加载完成！")

    cool_progress.start_spinner("正在解析HTML...")
    soup = BeautifulSoup(html, "html.parser")
    result: Dict[str, Any] = {
        "source_url": source_url,
        "desc": "", "author": "", "likes": "",
        "commentCount": "", "favs": "", "shares": "", "publishTime": "",
        "comments": [],
    }
    cool_progress.stop_spinner()

    # --- 1. 提取视频描述 ---
    cool_progress.cool_print("正在提取视频信息...", "info")
    desc_el = soup.find(attrs={"data-e2e": "detail-video-info"})
    if desc_el:
        h1 = desc_el.find("h1")
        if h1:
            result["desc"] = h1.get_text(strip=True)
        else:
            result["desc"] = desc_el.get_text(strip=True)

    # --- 2. 提取作者 ---
    user_info = soup.find(attrs={"data-e2e": "user-info"})
    if user_info:
        avatar_img = user_info.find("img")
        if avatar_img and avatar_img.get("alt"):
            result["author"] = avatar_img["alt"].strip().rstrip("头像")
    if not result["author"] and user_info:
        title_el = user_info.select_one('[data-click-from="title"]')
        if title_el:
            result["author"] = title_el.get_text(strip=True)

    # --- 3. 提取评论总数 ---
    comment_header = soup.find(class_=lambda x: x and "comment-header" in str(x))
    if comment_header:
        header_text = comment_header.get_text(strip=True)
        m = re.search(r"全部评论[\(\（](\d+)[\)\）]", header_text)
        if m:
            result["commentCount"] = m.group(1)

    # --- 4. 提取评论列表 ---
    cool_progress.cool_print("正在提取评论...", "star")
    comment_items = soup.find_all(attrs={"data-e2e": "comment-item"})
    total_comments = len(comment_items)
    seen = set()
    time_re = re.compile(r'^(刚刚|昨天|\d+年前|\d+月前|\d+天前|\d{4}-\d{2}-\d{2})$')

    for idx, item in enumerate(comment_items, 1):
        # 4a. 昵称
        name = ""
        name_link = item.select_one(".comment-item-info-wrap a")
        if name_link:
            name = _extract_deep_text(name_link).replace("...", "").strip()

        # 4b. 内容 — DOM层次结构
        content = ""
        info_wrap = item.select_one(".comment-item-info-wrap")
        if info_wrap and info_wrap.parent:
            content_area = info_wrap.parent
            found_info = False
            for child in content_area.children:
                if not hasattr(child, 'get_text'):
                    continue
                child_text = child.get_text(strip=True)
                if not child_text:
                    continue
                if child == info_wrap:
                    found_info = True
                    continue
                if not found_info:
                    continue
                if time_re.match(child_text):
                    break
                if re.match(r'^\d+$', child_text):
                    break
                if child_text in ('分享', '回复'):
                    break
                content = child_text
                break

        # 兜底
        if not content and info_wrap:
            for sibling in info_wrap.next_siblings:
                if not hasattr(sibling, 'get_text'):
                    continue
                t = sibling.get_text(strip=True)
                if not t:
                    continue
                if time_re.match(t) or t in ('分享', '回复') or re.match(r'^\d+$', t):
                    break
                content = t
                break

        # 4c. 时间
        time_val = ""
        for span in item.find_all("span"):
            t = span.get_text(strip=True)
            if time_re.match(t):
                time_val = t
                break

        # 4d. 点赞数
        likes = 0
        stats = item.select_one(".comment-item-stats-container")
        target = stats if stats else item
        for span in target.find_all("span"):
            t = span.get_text(strip=True)
            m = re.match(r'^(\d+)$', t)
            if m and len(t) < 7:
                likes = int(m.group(1))
                break

        # 4e. 回复数
        replies = 0
        reply_btn = item.select_one(".comment-reply-expand-btn")
        if reply_btn:
            t = reply_btn.get_text(strip=True)
            m = re.search(r'展开(\d+)条回复', t)
            if m:
                replies = int(m.group(1))

        # 4f. 作者赞过
        author_liked = False
        tag_el = item.select_one(".comment-item-tag")
        if tag_el and "作者赞过" in tag_el.get_text():
            author_liked = True

        # 清理内容
        content = content.replace("作者赞过", "").strip()

        # 去重
        key = (name + "|" + content[:50]).strip()
        if key and key not in seen and (len(content) > 1 or len(name) > 0):
            seen.add(key)
            result["comments"].append({
                "name": name, "content": content,
                "likes": likes, "time": time_val,
                "replies": replies, "authorLiked": author_liked,
            })
        
        # 显示进度
        if idx % 50 == 0 and total_comments > 0:
            cool_progress.print_progress_bar(idx, total_comments, prefix="解析评论中...")
    
    # 清除进度条行
    if total_comments >= 50:
        print()

    cool_progress.cool_print(f"视频描述: {result['desc'][:60] if result['desc'] else '(未获取)'}...", "info")
    cool_progress.cool_print(f"作者: {result['author'] or '(未获取)'}", "info")
    cool_progress.cool_print(f"评论数: {len(result['comments'])}", "success")
    return result


# ========== Playwright 爬取 ==========
def scrape_with_playwright(
    url: str, max_comments: int = MAX_COMMENTS, headless: bool = False,
    login_only: bool = False,
) -> Dict[str, Any]:
    """使用Playwright爬取抖音视频信息 + 评论"""
    # 智能决定是否 headless（有缓存 session 则 headless，否则弹浏览器登录）
    headless = get_smart_headless("douyin", headless)
    cool_progress.cool_print(f"启动浏览器（headless={headless}）", "rocket")
    cool_progress.cool_print(f"目标URL: {url}", "info")

    result: Dict[str, Any] = {
        "source_url": url,
        "desc": "", "author": "", "likes": "",
        "commentCount": "", "favs": "", "shares": "", "publishTime": "",
        "comments": [],
    }

    with sync_playwright() as p:
        import subprocess
        r = subprocess.run(
            [sys.executable, "-m", "playwright", "install", "--dry-run", "chromium"],
            capture_output=True, text=True,
        )
        if "already" not in r.stdout and "already" not in r.stderr:
            cool_progress.cool_print("正在安装 Playwright Chromium...", "warning")
            cool_progress.start_spinner("正在安装浏览器...")
            subprocess.run([sys.executable, "-m", "playwright", "install", "chromium"], check=True)
            cool_progress.stop_spinner("浏览器安装完成！")

        context = create_browser_context(p, headless=headless)
        page = context.new_page()
        page.add_init_script(ANTI_DETECTION_JS)

        try:
            # ===== 登录检查 =====
            logged_in = ensure_session(page, "douyin")
            if not logged_in:
                cool_progress.cool_print("未登录，部分功能可能受限", "warning")

            # login-only 模式：登录完成就退出
            if login_only:
                cool_progress.cool_print("登录模式完成", "success")
                context.close()
                return result
            cool_progress.print_step(1, 5, "加载抖音页面", "正在加载目标视频页面...")
            # 重试加载
            max_retries = 3
            for attempt in range(max_retries):
                try:
                    cool_progress.start_spinner(f"正在加载页面... (尝试 {attempt + 1}/{max_retries})")
                    page.goto(url, timeout=90000, wait_until="domcontentloaded")
                    page.wait_for_timeout(5000)
                    cool_progress.stop_spinner("页面加载成功！")
                    break
                except Exception as e:
                    cool_progress.stop_spinner()
                    cool_progress.cool_print(f"加载失败: {e}", "warning")
                    if attempt < max_retries - 1:
                        cool_progress.cool_print(f"{5 * (attempt + 1)}秒后重试...", "warning")
                        time.sleep(5 * (attempt + 1))
                    else:
                        raise Exception(f"页面加载失败（已重试{max_retries}次）: {e}")

            cool_progress.print_step(2, 5, "等待评论区加载", "正在定位评论内容...")
            # 等待评论区
            cool_progress.start_spinner("等待评论区加载...")
            try:
                page.wait_for_selector('[data-e2e="comment-list"]', timeout=15000)
                cool_progress.stop_spinner("评论区容器加载成功！")
            except Exception:
                cool_progress.stop_spinner()
                cool_progress.cool_print("未检测到评论区容器", "warning")

            try:
                page.wait_for_selector('[data-e2e="comment-item"]', timeout=10000)
                cool_progress.cool_print("评论项已出现", "success")
            except Exception:
                cool_progress.cool_print("未检测到评论项", "warning")

            page.wait_for_timeout(3000)

            cool_progress.print_step(3, 5, "提取视频信息", "正在获取视频详情...")
            # 提取视频信息
            video_info = page.evaluate(EXTRACT_VIDEO_INFO_JS)
            for key in ["desc", "author", "followers", "totalLikes", "likes", "commentCount", "favs", "shares", "publishTime"]:
                if key in video_info and video_info[key]:
                    result[key] = video_info[key]
            
            # 炫酷显示视频信息
            cool_progress.cool_print(f"视频描述: {result['desc'][:60] if result['desc'] else '(未获取)'}...", "info")
            cool_progress.cool_print(f"作者: {result['author'] or '(未获取)'}", "info")
            cool_progress.cool_print(f"粉丝: {result.get('followers') or '(未获取)'}, 获赞: {result.get('totalLikes') or '(未获取)'}", "info")
            cool_progress.cool_print(f"点赞: {result.get('likes') or '(未获取)'}, 评论: {result.get('commentCount') or '(未获取)'}, 收藏: {result.get('favs') or '(未获取)'}, 转发: {result.get('shares') or '(未获取)'}", "info")

            COUNT_COMMENTS_JS = """
            () => {
                return document.querySelectorAll('[data-e2e="comment-item"]').length;
            }
            """

            WHEEL_COMMENT_CONTAINER_JS = """
            () => {
                const containers = [
                    document.querySelector('[data-e2e="comment-list"]'),
                    document.querySelector('.comment-mainContent'),
                ];
                for (const el of containers) {
                    if (el) {
                        el.dispatchEvent(new WheelEvent('wheel', {
                            deltaY: 800, deltaX: 0,
                            bubbles: true, cancelable: true,
                        }));
                        return { target: el.tagName, className: el.className };
                    }
                }
                window.dispatchEvent(new WheelEvent('wheel', {
                    deltaY: 800, deltaX: 0,
                    bubbles: true, cancelable: true,
                }));
                return { target: 'window' };
            }
            """

            cool_progress.print_step(4, 5, "滚动加载评论", "正在获取所有评论内容...")
            cool_progress.cool_print("开始滚动加载评论...", "star")

            page.evaluate("""
            () => {
                const el = document.querySelector('[data-e2e="comment-list"]');
                if (el) el.scrollIntoView({ behavior: 'instant', block: 'start' });
            }
            """)
            page.wait_for_timeout(2000)

            last_count = 0
            no_change = 0

            for rnd in range(SCROLL_ROUNDS):
                page.evaluate(WHEEL_COMMENT_CONTAINER_JS)
                page.wait_for_timeout(SCROLL_WAIT_MS)

                cur = page.evaluate(COUNT_COMMENTS_JS)
                if cur == last_count:
                    no_change += 1
                    if no_change >= NO_CHANGE_THRESHOLD:
                        cool_progress.cool_print(f"连续{NO_CHANGE_THRESHOLD}轮评论数不变（共{cur}条），停止滚动", "success")
                        break
                else:
                    no_change = 0
                    cool_progress.print_progress_bar(min(cur, max_comments), max_comments, prefix=f"已加载 {cur} 条评论")
                last_count = cur

                if cur >= max_comments:
                    cool_progress.cool_print(f"已达上限 {max_comments} 条", "success")
                    break

            # 清除进度条的行
            print()

            cool_progress.cool_print("正在展开子评论...", "star")
            
            expand_buttons_clicked = 0
            for attempt in range(10):
                btn_count = page.locator('.comment-reply-expand-btn').count()
                if btn_count == 0:
                    cool_progress.cool_print(f"第{attempt+1}轮无展开按钮，停止", "info")
                    break
                
                for i in range(btn_count):
                    try:
                        btn = page.locator('.comment-reply-expand-btn').first
                        btn.scroll_into_view_if_needed(timeout=3000)
                        btn.click(timeout=3000)
                        expand_buttons_clicked += 1
                        page.wait_for_timeout(400)
                    except Exception:
                        pass
                
                cool_progress.cool_print(f"第{attempt+1}轮点击了 {btn_count} 个展开回复按钮", "info")
                
                page.wait_for_timeout(2000)
                
                # 点击"展开更多"按钮（加载更多子评论）- 多轮点击
                for more_attempt in range(20):
                    try:
                        more_btns = page.locator('text=展开更多')
                        more_count = more_btns.count()
                        if more_count == 0:
                            break
                        clicked_more = 0
                        for j in range(min(more_count, 5)):
                            try:
                                more_btns.nth(j).scroll_into_view_if_needed(timeout=2000)
                                more_btns.nth(j).click(timeout=2000)
                                expand_buttons_clicked += 1
                                clicked_more += 1
                                page.wait_for_timeout(1500)
                            except Exception:
                                pass
                        if clicked_more == 0:
                            break
                    except Exception:
                        break
                
                # 滚动评论区以加载子评论
                for scroll_round in range(8):
                    page.evaluate(WHEEL_COMMENT_CONTAINER_JS)
                    page.wait_for_timeout(800)
                
                # 检查是否还有未展开的
                remaining = page.locator('.comment-reply-expand-btn').count()
                remaining_more = page.locator('text=展开更多').count()
                if remaining == 0 and remaining_more == 0:
                    cool_progress.cool_print("所有子评论已展开", "success")
                    break
            
            cool_progress.cool_print(f"共点击了 {expand_buttons_clicked} 个展开按钮", "success")
            
            # 再次滚动确保子评论加载
            cool_progress.cool_print("再次滚动以加载子评论...", "star")
            for rnd in range(20):
                page.evaluate(WHEEL_COMMENT_CONTAINER_JS)
                page.wait_for_timeout(SCROLL_WAIT_MS)

            cool_progress.print_step(5, 5, "提取评论数据", "正在整理获取到的评论...")
            cool_progress.start_spinner("正在提取评论...")
            page.wait_for_timeout(3000)

            raw_comments = page.evaluate(EXTRACT_COMMENTS_JS)
            all_comments = []
            seen_keys = set()
            for c in raw_comments:
                key = (c.get("name", "") + "|" + c.get("content", "")[:50]).strip()
                if key and key not in seen_keys:
                    seen_keys.add(key)
                    all_comments.append(c)
                    if len(all_comments) >= max_comments:
                        break

            result["comments"] = all_comments
            cool_progress.stop_spinner(f"成功提取 {len(all_comments)} 条评论！")

        except Exception as e:
            cool_progress.stop_spinner()
            cool_progress.cool_print(f"爬取过程出错: {e}", "error")
            import traceback
            traceback.print_exc()
        finally:
            context.close()
            cool_progress.cool_print("浏览器已关闭", "info")

    return result


# ========== Excel 导出 ==========
def save_to_excel(results: List[Dict[str, Any]], filepath: str):
    """将抖音数据保存为Excel（单表：视频信息 + 评论）"""
    cool_progress.cool_print(f"正在生成Excel: {filepath}", "rocket")
    cool_progress.start_spinner("正在生成Excel文件...")

    wb = Workbook()
    hf = Font(name="Microsoft YaHei", bold=True, size=11, color="FFFFFF")
    hfl = PatternFill(start_color="FF2F5496", end_color="FF2F5496", fill_type="solid")
    ha = Alignment(horizontal="center", vertical="center", wrap_text=True)
    df = Font(name="Microsoft YaHei", size=10)
    da = Alignment(vertical="top", wrap_text=True)
    ca = Alignment(horizontal="center", vertical="center", wrap_text=True)
    bd = Border(
        left=Side("thin", "FFD0D0D0"), right=Side("thin", "FFD0D0D0"),
        top=Side("thin", "FFD0D0D0"), bottom=Side("thin", "FFD0D0D0"),
    )

    ws = wb.active
    ws.title = "抖音数据"

    headers = [
        "序号", "来源链接",
        "视频作者", "粉丝数", "获赞数", "视频描述/文案", "点赞数", "评论总数",
        "收藏数", "转发数", "发布时间",
        "评论者昵称", "评论内容", "评论点赞", "评论时间",
        "回复数", "作者赞过", "是否子评论",
    ]
    widths = [6, 40, 16, 10, 10, 60, 10, 10, 10, 10, 14, 16, 60, 10, 14, 8, 8, 10]

    total_rows = sum(1 if not r.get("comments", []) else len(r.get("comments", [])) for r in results)
    current_row = 0

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font, c.fill, c.alignment, c.border = hf, hfl, ha, bd
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col)].width = w

    row_idx = 2
    seq = 0
    for r in results:
        comments = r.get("comments", [])
        if not comments:
            seq += 1
            row_data = [
                seq,
                r.get("source_url", ""),
                r.get("author", ""), r.get("followers", ""), r.get("totalLikes", ""),
                r.get("desc", ""),
                r.get("likes", ""), r.get("commentCount", ""),
                r.get("favs", ""), r.get("shares", ""), r.get("publishTime", ""),
                "", "", "", "", "", "", "",
            ]
            for col, val in enumerate(row_data, 1):
                c = ws.cell(row=row_idx, column=col, value=val)
                c.font, c.border = df, bd
                c.alignment = da if col in (6, 13) else ca
            row_idx += 1
            current_row += 1
        else:
            for cm in comments:
                seq += 1
                row_data = [
                    seq,
                    r.get("source_url", ""),
                    r.get("author", ""), r.get("followers", ""), r.get("totalLikes", ""),
                    r.get("desc", ""),
                    r.get("likes", ""), r.get("commentCount", ""),
                    r.get("favs", ""), r.get("shares", ""), r.get("publishTime", ""),
                    cm.get("name", ""), cm.get("content", ""),
                    cm.get("likes", 0), cm.get("time", ""),
                    cm.get("replies", 0),
                    "是" if cm.get("authorLiked") else "",
                    "是" if cm.get("isChild") else "",
                ]
                for col, val in enumerate(row_data, 1):
                    c = ws.cell(row=row_idx, column=col, value=val)
                    c.font, c.border = df, bd
                    c.alignment = da if col in (6, 13) else ca
                row_idx += 1
                current_row += 1
                if current_row % 100 == 0:
                    cool_progress.print_progress_bar(current_row, total_rows, prefix="写入Excel中...")

    if row_idx > 2:
        ws.freeze_panes = "A2"

    wb.save(filepath)
    cool_progress.stop_spinner(f"Excel文件已成功保存到: {filepath}")
    return filepath


# ========== 批量URL列表 ==========
def load_urls_from_file(filepath: str) -> List[str]:
    urls = []
    with open(filepath, "r", encoding="utf-8") as f:
        for line in f:
            url = line.strip()
            if url and not url.startswith("#"):
                urls.append(url)
    return urls


# ========== 主程序 ==========
def main():
    args = parse_args()

    # 炫酷启动动画
    cool_progress.print_banner()

    urls_to_scrape = []

    # login-only 模式：只登录不爬取
    if args.login_only:
        cool_progress.cool_print("🔑 登录模式：仅打开浏览器登录抖音", "rocket")
        result = scrape_with_playwright(
            args.url, max_comments=0, headless=False, login_only=True,
        )
        return

    if args.urls_file:
        urls_to_scrape = load_urls_from_file(args.urls_file)
        cool_progress.cool_print(f"模式: 批量爬取（{len(urls_to_scrape)}个链接）", "rocket")
    elif args.urls:
        urls_to_scrape = load_urls_from_file(args.urls)
        cool_progress.cool_print(f"模式: 批量爬取（{len(urls_to_scrape)}个链接）", "rocket")
    elif args.local_file:
        cool_progress.cool_print("模式: 本地HTML解析", "star")
        cool_progress.cool_print(f"文件: {args.local_file}", "info")
        print()
        result = parse_local_html(args.local_file, args.url)
        output_path = args.output_file or os.path.join(
            OUTPUT_DIR, f"douyin_local_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        save_to_excel([result], output_path)
        cool_final_msg = f"\n{TC.BRIGHT_GREEN}{TC.BOLD}╔════════════════════════════════════════════════════════════════╗{TC.RESET}"
        cool_final_msg += f"\n{TC.BRIGHT_GREEN}{TC.BOLD}║{TC.RESET}  {TC.BRIGHT_CYAN}{TC.BOLD}✅  解 析 完 成  ✅{TC.RESET}{' ' * 40}{TC.BRIGHT_GREEN}{TC.BOLD}║{TC.RESET}"
        cool_final_msg += f"\n{TC.BRIGHT_GREEN}{TC.BOLD}╚════════════════════════════════════════════════════════════════╝{TC.RESET}"
        print(cool_final_msg)
        cool_progress.cool_print(f"视频描述: {result['desc'][:50]}...", "success")
        cool_progress.cool_print(f"评论数: {len(result['comments'])}", "info")
        print()
        cool_progress.cool_print(f"输出文件: {output_path}", "rocket")
        print()
        return
    else:
        urls_to_scrape = [args.url]
        cool_progress.cool_print("模式: 单链接爬取", "star")
        cool_progress.cool_print(f"URL: {args.url}", "info")
        cool_progress.cool_print(f"最多评论: {args.max_comments}", "info")
        print()

    # 爬取
    all_results = []
    for idx, url in enumerate(urls_to_scrape, 1):
        if len(urls_to_scrape) > 1:
            cool_progress.print_step(idx, len(urls_to_scrape), f"正在处理视频 {idx}/{len(urls_to_scrape)}", url)
        result = scrape_with_playwright(url, max_comments=args.max_comments, headless=args.headless, login_only=False)
        all_results.append(result)
        if len(urls_to_scrape) > 1:
            cool_progress.cool_print(f"累计: {len(all_results)}/{len(urls_to_scrape)} 个视频", "info")

    if not any(r.get("desc") or r.get("comments") for r in all_results):
        log("⚠️ 未获取到数据", level="warning")
        cool_progress.cool_print("未获取到数据，请检查URL或是否需要登录", "warning")
        return

    output_path = args.output_file or os.path.join(
        OUTPUT_DIR, f"douyin_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    )
    save_to_excel(all_results, output_path)

    # 炫酷完成总结
    cool_final_msg = f"\n{TC.BRIGHT_MAGENTA}{TC.BOLD}╔════════════════════════════════════════════════════════════════╗{TC.RESET}"
    cool_final_msg += f"\n{TC.BRIGHT_MAGENTA}{TC.BOLD}║{TC.RESET}  {TC.BRIGHT_GREEN}{TC.BOLD}🎉  爬 取 完 成  🎉{TC.RESET}{' ' * 40}{TC.BRIGHT_MAGENTA}{TC.BOLD}║{TC.RESET}"
    cool_final_msg += f"\n{TC.BRIGHT_MAGENTA}{TC.BOLD}╚════════════════════════════════════════════════════════════════╝{TC.RESET}"
    print(cool_final_msg)
    for r in all_results:
        author = r.get('author', '?')
        desc = r.get('desc', '?')[:40]
        likes = r.get('likes', '?')
        comment_count = r.get('commentCount', '?')
        comments_num = len(r.get('comments', []))
        print(f"\n  {TC.BRIGHT_CYAN}📹 {author}{TC.RESET} - {TC.BRIGHT_YELLOW}{desc}{TC.RESET}...")
        print(f"     {TC.BRIGHT_GREEN}👍 {likes}{TC.RESET}  {TC.BRIGHT_BLUE}💬 {comment_count}{TC.RESET}  {TC.BRIGHT_MAGENTA}🗣 {comments_num}条评论{TC.RESET}")
    print(f"\n  {TC.BRIGHT_YELLOW}📊 输出文件: {output_path}{TC.RESET}")
    print()


if __name__ == "__main__":
    main()
